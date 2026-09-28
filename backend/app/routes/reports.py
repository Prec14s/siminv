import io
from datetime import date, datetime
from flask import Blueprint, request, send_file
from sqlalchemy import func
from ..extensions import db
from ..models import Item, StockTransaction, StockTransactionItem, Request, Setting, User
from ..utils import ok, roles_required, parse_date, ApiError, MANAGERS

bp = Blueprint("reports", __name__)

TRX_LABEL = {"initial": "Stok Awal", "in": "Masuk", "out": "Keluar", "adjustment": "Penyesuaian",
             "return": "Pengembalian", "reversal": "Pembatalan"}
REQ_STATUS = {"pending": "Menunggu", "approved": "Disetujui", "rejected": "Ditolak",
              "handed_over": "Dipinjam", "returned": "Dikembalikan", "completed": "Selesai",
              "cancelled": "Dibatalkan"}


def period():
    d_from = parse_date(request.args.get("date_from"), required=False)
    d_to = parse_date(request.args.get("date_to"), required=False)
    label = f"{d_from or 'awal'} s.d. {d_to or 'sekarang'}"
    return d_from, d_to, label


def report_stock():
    as_of = parse_date(request.args.get("as_of"), required=False)
    q = Item.query.filter(Item.is_deleted.is_(False))
    if cid := request.args.get("category_id", type=int):
        q = q.filter(Item.category_id == cid)
    if lid := request.args.get("location_id", type=int):
        q = q.filter(Item.location_id == lid)
    items = q.order_by(Item.code).all()
    later = {}
    if as_of:
        later = dict(db.session.query(StockTransactionItem.item_id, func.sum(StockTransactionItem.quantity))
                     .join(StockTransaction).filter(StockTransaction.trx_date > as_of)
                     .group_by(StockTransactionItem.item_id).all())
    rows, total = [], 0.0
    for i in items:
        stock = i.stock - int(later.get(i.id) or 0)
        value = stock * float(i.price)
        total += value
        rows.append({"code": i.code, "name": i.name, "category": i.category.name,
                     "location": i.location.name, "stock": stock, "unit": i.unit.name,
                     "min_stock": i.min_stock, "price": float(i.price), "value": value,
                     "status": "Habis" if stock <= 0 else ("Menipis" if stock <= i.min_stock else "Aman")})
    cols = [("code", "Kode"), ("name", "Nama Barang"), ("category", "Kategori"), ("location", "Lokasi"),
            ("stock", "Stok"), ("unit", "Satuan"), ("min_stock", "Min"), ("price", "Harga"),
            ("value", "Nilai"), ("status", "Status")]
    return {"title": "Laporan Stok Barang", "subtitle": f"Per tanggal {as_of or date.today()}",
            "columns": cols, "rows": rows,
            "summary": {"Jumlah jenis barang": len(rows), "Total nilai inventaris": total}}


def report_transactions():
    d_from, d_to, label = period()
    q = (db.session.query(StockTransactionItem, StockTransaction).join(StockTransaction))
    if d_from:
        q = q.filter(StockTransaction.trx_date >= d_from)
    if d_to:
        q = q.filter(StockTransaction.trx_date <= d_to)
    if t := request.args.get("type"):
        q = q.filter(StockTransaction.type.in_(t.split(",")))
    rows = []
    t_in = t_out = 0
    for d, t in q.order_by(StockTransaction.trx_date, StockTransaction.id).all():
        rows.append({"date": t.trx_date.isoformat(), "number": t.trx_number, "type": TRX_LABEL[t.type],
                     "code": d.item.code, "name": d.item.name,
                     "in": d.quantity if d.quantity > 0 else 0, "out": -d.quantity if d.quantity < 0 else 0,
                     "party": (t.supplier.name if t.supplier else t.recipient) or "-",
                     "user": t.creator.name, "notes": d.reason or t.notes or ""})
        t_in += max(d.quantity, 0)
        t_out += max(-d.quantity, 0)
    cols = [("date", "Tanggal"), ("number", "No. Transaksi"), ("type", "Tipe"), ("code", "Kode"),
            ("name", "Barang"), ("in", "Masuk"), ("out", "Keluar"), ("party", "Supplier/Penerima"),
            ("user", "Petugas"), ("notes", "Keterangan")]
    return {"title": "Laporan Transaksi Barang", "subtitle": f"Periode {label}", "columns": cols,
            "rows": rows, "summary": {"Total masuk": t_in, "Total keluar": t_out}}


def report_requests():
    d_from, d_to, label = period()
    q = Request.query.join(User, User.id == Request.requester_id)
    if d_from:
        q = q.filter(func.date(Request.created_at) >= d_from)
    if d_to:
        q = q.filter(func.date(Request.created_at) <= d_to)
    if t := request.args.get("type"):
        q = q.filter(Request.type == t)
    if s := request.args.get("status"):
        q = q.filter(Request.status == s)
    if div := request.args.get("division"):
        q = q.filter(User.division.ilike(f"%{div}%"))
    rows = []
    for r in q.order_by(Request.created_at).all():
        for ri in r.items:
            rows.append({"date": r.created_at.date().isoformat(), "number": r.request_number,
                         "type": "Peminjaman" if r.type == "loan" else "Permintaan",
                         "requester": r.requester.name, "division": r.requester.division or "-",
                         "item": ri.item.name, "requested": ri.qty_requested,
                         "approved": ri.qty_approved if ri.qty_approved is not None else "-",
                         "status": REQ_STATUS[r.status] + (" (terlambat)" if r.is_overdue else "")})
    cols = [("date", "Tanggal"), ("number", "No. Pengajuan"), ("type", "Jenis"), ("requester", "Pengaju"),
            ("division", "Divisi"), ("item", "Barang"), ("requested", "Diminta"),
            ("approved", "Disetujui"), ("status", "Status")]
    return {"title": "Laporan Permintaan & Peminjaman", "subtitle": f"Periode {label}",
            "columns": cols, "rows": rows,
            "summary": {"Jumlah pengajuan": len({r['number'] for r in rows})}}


def report_stock_card():
    iid = request.args.get("item_id", type=int)
    if not iid:
        raise ApiError("Pilih barang untuk kartu stok", 422)
    item = Item.query.get_or_404(iid)
    d_from, d_to, label = period()
    base = (db.session.query(StockTransactionItem, StockTransaction).join(StockTransaction)
            .filter(StockTransactionItem.item_id == iid))
    opening = 0
    if d_from:
        prev = (base.filter(StockTransaction.trx_date < d_from)
                .order_by(StockTransaction.trx_date.desc(), StockTransactionItem.id.desc()).first())
        opening = prev[0].stock_after if prev else 0
    q = base
    if d_from:
        q = q.filter(StockTransaction.trx_date >= d_from)
    if d_to:
        q = q.filter(StockTransaction.trx_date <= d_to)
    rows = []
    balance = opening
    t_in = t_out = 0
    for d, t in q.order_by(StockTransaction.trx_date, StockTransactionItem.id).all():
        balance += d.quantity
        t_in += max(d.quantity, 0)
        t_out += max(-d.quantity, 0)
        rows.append({"date": t.trx_date.isoformat(), "number": t.trx_number, "type": TRX_LABEL[t.type],
                     "in": max(d.quantity, 0), "out": max(-d.quantity, 0), "balance": balance,
                     "notes": d.reason or t.notes or ""})
    cols = [("date", "Tanggal"), ("number", "No. Transaksi"), ("type", "Tipe"), ("in", "Masuk"),
            ("out", "Keluar"), ("balance", "Saldo"), ("notes", "Keterangan")]
    return {"title": f"Kartu Stok: {item.code} - {item.name}", "subtitle": f"Periode {label}",
            "columns": cols, "rows": rows,
            "summary": {"Saldo awal": opening, "Total masuk": t_in, "Total keluar": t_out,
                        "Saldo akhir": balance}}


REPORTS = {"stock": report_stock, "transactions": report_transactions,
           "requests": report_requests, "stock-card": report_stock_card}


@bp.get("/reports/<kind>")
@roles_required(*MANAGERS)
def get_report(kind):
    if kind not in REPORTS:
        raise ApiError("Jenis laporan tidak dikenal", 404)
    rep = REPORTS[kind]()
    fmt = request.args.get("format", "json")
    s = db.session.get(Setting, "institution_name")
    rep["institution"] = s.value if s else "Instansi Saya"
    filename = f"laporan-{kind}-{datetime.now():%Y%m%d-%H%M}"
    if fmt == "xlsx":
        return send_file(to_xlsx(rep), as_attachment=True, download_name=f"{filename}.xlsx",
                         mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    if fmt == "pdf":
        return send_file(to_pdf(rep), as_attachment=True, download_name=f"{filename}.pdf",
                         mimetype="application/pdf")
    rep["columns"] = [{"key": k, "label": l} for k, l in rep["columns"]]
    return ok(rep)


def to_xlsx(rep):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = "Laporan"
    ws.append([rep["institution"]])
    ws.append([rep["title"]])
    ws.append([rep["subtitle"]])
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"].font = Font(bold=True, size=14)
    ws.append([])
    ws.append([l for _, l in rep["columns"]])
    head_row = ws.max_row
    for c in ws[head_row]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3A4D")
        c.alignment = Alignment(horizontal="center")
    for r in rep["rows"]:
        ws.append([r[k] for k, _ in rep["columns"]])
    ws.append([])
    for k, v in (rep.get("summary") or {}).items():
        ws.append([k, v])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
    for idx, (k, l) in enumerate(rep["columns"], 1):
        width = max([len(str(l))] + [len(str(r[k])) for r in rep["rows"]] or [10])
        ws.column_dimensions[get_column_letter(idx)].width = min(max(width + 2, 8), 45)
    ws.freeze_panes = ws.cell(head_row + 1, 1)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def to_pdf(rep):
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm,
                            topMargin=12 * mm, bottomMargin=12 * mm, title=rep["title"])
    styles = getSampleStyleSheet()
    cell = ParagraphStyle("cell", parent=styles["Normal"], fontSize=7.5, leading=9)
    head = ParagraphStyle("head", parent=cell, textColor=colors.white, fontName="Helvetica-Bold")

    def fmt(v):
        if isinstance(v, float):
            return f"{v:,.0f}".replace(",", ".")
        return str(v)

    story = [Paragraph(rep["institution"], styles["Normal"]),
             Paragraph(rep["title"], styles["Title"]),
             Paragraph(rep["subtitle"], styles["Normal"]), Spacer(1, 6 * mm)]
    data = [[Paragraph(l, head) for _, l in rep["columns"]]]
    for r in rep["rows"]:
        data.append([Paragraph(fmt(r[k]), cell) for k, _ in rep["columns"]])
    if len(data) == 1:
        data.append([Paragraph("Tidak ada data", cell)] + [""] * (len(rep["columns"]) - 1))
    t = Table(data, repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F3A4D")),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#C9CED3")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F5F6")]),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    story.append(t)
    if rep.get("summary"):
        story.append(Spacer(1, 5 * mm))
        for k, v in rep["summary"].items():
            story.append(Paragraph(f"<b>{k}:</b> {fmt(v)}", styles["Normal"]))
    story.append(Spacer(1, 4 * mm))
    story.append(Paragraph(f"Dicetak {datetime.now():%d-%m-%Y %H:%M}", cell))
    doc.build(story)
    buf.seek(0)
    return buf
